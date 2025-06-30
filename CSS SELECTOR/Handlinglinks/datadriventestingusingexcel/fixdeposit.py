# import os
# import time
# import openpyxl
# from selenium import webdriver
# from selenium.webdriver.common.by import By
# driver=webdriver.Chrome()
# driver.maximize_window()
# time.sleep(2)
# driver.get("https://www.moneycontrol.com/fixed-income/calculator/state-bank-of-india-sbi/fixed-deposit-calculator-SBI-BSB001.html")
# file=os.getcwd()+r"\caldata.xlsx"
# print(file)
# workbook=openpyxl.load_workbook(file)
# sheet=workbook["Sheet1"]
#
# rows=sheet.max_row
# print("The no.of rows are",rows)
# columns=sheet.max_column
# print("The no. of columns are",columns)
# time.sleep(2)
#
# for r in range(1,rows+1):
#     for c in range(1,columns+1):
#         print(sheet.cell(r,c).value,end="   ")
#     print()

import os
import time
import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.select import Select

import XLUtils
driver=webdriver.Chrome()
driver.maximize_window()
time.sleep(2)
driver.get("https://www.moneycontrol.com/fixed-income/calculator/state-bank-of-india-sbi/fixed-deposit-calculator-SBI-BSB001.html")
file=os.getcwd()+r"\caldata.xlsx"
rows=XLUtils.getRowCount(file,"Sheet1")
# print(rows)

for r in range(2,rows+1):
    pric=XLUtils.readData(file,"Sheet1",r,1)
    rateofinterest=XLUtils.readData(file,"Sheet1",r,2)
    per1=XLUtils.readData(file,"Sheet1",r,3)
    per2=XLUtils.readData(file,"Sheet1",r,4)
    fre=XLUtils.readData(file,"Sheet1",r,5)
    exp_mvalue=XLUtils.readData(file,"Sheet1",r,6)

    driver.find_element(By.XPATH,'//*[@id="principal"]').send_keys(pric)
    driver.find_element(By.XPATH,'//*[@id="interest"]').send_keys(rateofinterest)
    driver.find_element(By.XPATH,'//*[@id="tenure"]').send_keys(per1)
    perioddrp=Select(driver.find_element(By.XPATH,'//*[@id="tenurePeriod"]'))
    perioddrp.select_by_visible_text(per2)
    frequencydrp=Select(driver.find_element(By.XPATH,'//*[@id="frequency"]'))
    frequencydrp.select_by_visible_text(fre)
    time.sleep(4)

    click_button=driver.find_element(By.XPATH,'//*[@id="fdMatVal"]/div[2]/a[1]/img')
    driver.execute_script("arguments[0].click();",click_button)

    time.sleep(1)

    act_value = driver.find_element(By.XPATH, '//*[@id="resp_matval"]').text

    if float(exp_mvalue)==float(act_value):
        print("test passed")
        XLUtils.writeData(file,"Sheet1",r,8,"Passed")
        XLUtils.fillGreenColor(file,"Sheet1",r,8)
    else:
        print("Test failed")
        XLUtils.writeData(file,"Sheet1",r,8,"Failed")
        XLUtils.fillRedColor(file,"Sheet1",r,8)

    clear_Button=driver.find_element(By.XPATH,'//*[@id="fdMatVal"]/div[2]/a[2]/img')
    driver.execute_script("arguments[0].click() ;",clear_Button)

# for r in range(2, rows + 1):
#     # Read Excel data
#     pric = XLUtils.readData(file, "Sheet1", r, 1)
#     rateofinterest = XLUtils.readData(file, "Sheet1", r, 2)
#     per1 = XLUtils.readData(file, "Sheet1", r, 3)
#     per2 = XLUtils.readData(file, "Sheet1", r, 4)
#     fre = XLUtils.readData(file, "Sheet1", r, 5)
#     exp_mvalue = XLUtils.readData(file, "Sheet1", r, 6)
#
#     # Fill form after clearing fields
#     driver.find_element(By.XPATH, '//*[@id="principal"]').clear()
#     driver.find_element(By.XPATH, '//*[@id="principal"]').send_keys(pric)
#
#     driver.find_element(By.XPATH, '//*[@id="interest"]').clear()
#     driver.find_element(By.XPATH, '//*[@id="interest"]').send_keys(rateofinterest)
#
#     driver.find_element(By.XPATH, '//*[@id="tenure"]').clear()
#     driver.find_element(By.XPATH, '//*[@id="tenure"]').send_keys(per1)
#
#     Select(driver.find_element(By.XPATH, '//*[@id="tenurePeriod"]')).select_by_visible_text(per2)
#     Select(driver.find_element(By.XPATH, '//*[@id="frequency"]')).select_by_visible_text(fre)
#
#     # Click calculate and read result
#     click_button = driver.find_element(By.XPATH, '//*[@id="fdMatVal"]/div[2]/a[1]/img')
#     driver.execute_script("arguments[0].click()", click_button)
#     time.sleep(3)
#
#     act_value = driver.find_element(By.XPATH, '//*[@id="resp_matval"]').text
#     print(f"Row {r} - Actual Maturity Value: {act_value}")




